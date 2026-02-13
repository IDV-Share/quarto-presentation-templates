import argparse
import json
import os
import re
import shutil
import sys
import time
import zipfile
from urllib.parse import urlparse


PLACEHOLDER_PREFIX = "VIDEO::"
PICTURE_TYPES = {11, 13, 14}
GROUP_TYPE = 6
FILL_PICTURE = 6
VISUAL_TYPES = {3, 7, 10, 11, 13, 14, 16}
COLUMN_PREFIX = "COLUMN::"
BOX_PREFIX = "BOX::"
BOX_END_PREFIX = "BOXEND::"
CAPTION_PREFIX = "CAPTION::"
LAYOUT_PREFIX = "LAYOUT::"
ESCAPABLE_MARKER_PREFIXES = (
    PLACEHOLDER_PREFIX,
    COLUMN_PREFIX,
    BOX_PREFIX,
    BOX_END_PREFIX,
    CAPTION_PREFIX,
    LAYOUT_PREFIX,
)
TEXTBOX_ORIENTATION = 1
EMU_PER_POINT = 12700.0


def truthy(value):
    return str(value).strip().lower() in {"1", "true", "yes", "on"}


def load_mapping(path):
    with open(path, "r", encoding="utf-8") as handle:
        data = json.load(handle)
    if not isinstance(data, list):
        raise ValueError("Mapping file must contain a list of video entries.")
    return data


def load_layout(path):
    with open(path, "r", encoding="utf-8") as handle:
        data = json.load(handle)
    if not isinstance(data, dict):
        return {"columns": [], "boxes": [], "captions": [], "caption_defaults": {}, "footer": "", "location": ""}
    columns = data.get("columns", [])
    boxes = data.get("boxes", [])
    captions = data.get("captions", [])
    slide_layouts = data.get("slide_layouts", [])
    caption_defaults = data.get("caption_defaults", {})
    if not isinstance(columns, list):
        columns = []
    if not isinstance(boxes, list):
        boxes = []
    if not isinstance(captions, list):
        captions = []
    if not isinstance(slide_layouts, list):
        slide_layouts = []
    if not isinstance(caption_defaults, dict):
        caption_defaults = {}
    return {
        "columns": columns,
        "boxes": boxes,
        "captions": captions,
        "slide_layouts": slide_layouts,
        "caption_defaults": caption_defaults,
        "footer": data.get("footer", "") or "",
        "location": data.get("location", "") or "",
    }


def resolve_path(base_dir, path_value):
    if not path_value:
        return ""
    if is_url(path_value):
        return path_value
    if os.path.isabs(path_value):
        return path_value
    return os.path.abspath(os.path.join(base_dir, path_value))


def collect_mapping(entries, base_dir):
    mapping = {}
    for entry in entries:
        if not isinstance(entry, dict):
            continue
        video_id = entry.get("id")
        if not video_id:
            continue
        mapping[video_id] = {
            "video": resolve_path(base_dir, entry.get("video", "")),
            "poster": resolve_path(base_dir, entry.get("poster", "")),
            "width": entry.get("width", "") or "",
            "height": entry.get("height", "") or "",
            "x": entry.get("x", "") or "",
            "y": entry.get("y", "") or "",
        }
    return mapping


def collect_layout(entries):
    columns = {}
    boxes = {}
    captions = []
    for entry in entries.get("columns", []):
        if isinstance(entry, dict) and entry.get("id"):
            columns[entry["id"]] = entry
    for entry in entries.get("boxes", []):
        if isinstance(entry, dict) and entry.get("id"):
            boxes[entry["id"]] = entry
    for entry in entries.get("captions", []):
        if isinstance(entry, dict):
            captions.append(entry)
    return columns, boxes, captions


def iter_shapes(slide):
    count = slide.Shapes.Count
    for index in range(count, 0, -1):
        yield slide.Shapes(index)


def is_picture(shape):
    try:
        return shape.Type in PICTURE_TYPES
    except Exception:
        return False


def is_group(shape):
    try:
        return shape.Type == GROUP_TYPE
    except Exception:
        return False


def has_text(shape):
    try:
        if shape.HasTextFrame and shape.TextFrame.HasText:
            return bool((shape.TextFrame.TextRange.Text or "").strip())
    except Exception:
        return False
    return False


def is_picture_fill(shape):
    try:
        return shape.Fill.Type == FILL_PICTURE
    except Exception:
        return False


def is_visual_shape(shape):
    try:
        shape_type = shape.Type
    except Exception:
        return False

    if shape_type in VISUAL_TYPES:
        return True

    if is_picture_fill(shape):
        return True

    if shape_type == GROUP_TYPE:
        try:
            count = shape.GroupItems.Count
        except Exception:
            return False
        for index in range(1, count + 1):
            try:
                item = shape.GroupItems(index)
            except Exception:
                continue
            if is_visual_shape(item):
                return True
    return False


def shape_bounds(shape):
    return (shape.Left, shape.Top, shape.Width, shape.Height)


def shape_alt_text(shape):
    try:
        return shape.AlternativeText or ""
    except Exception:
        return ""


def shape_text(shape):
    try:
        if shape.HasTextFrame and shape.TextFrame.HasText:
            return shape.TextFrame.TextRange.Text or ""
    except Exception:
        return ""
    return ""


def find_nearest_picture(slide, marker_shape):
    try:
        marker_cx = marker_shape.Left + (marker_shape.Width / 2)
        marker_cy = marker_shape.Top + (marker_shape.Height / 2)
    except Exception:
        return None

    candidates = []
    count = slide.Shapes.Count
    for index in range(1, count + 1):
        candidate = slide.Shapes(index)
        candidates.extend(picture_candidates(candidate))

    if not candidates:
        return None

    best = None
    best_dist = None
    for candidate in candidates:
        try:
            cx = candidate.Left + (candidate.Width / 2)
            cy = candidate.Top + (candidate.Height / 2)
        except Exception:
            continue
        dist = (cx - marker_cx) ** 2 + (cy - marker_cy) ** 2
        if best is None or dist < best_dist:
            best = candidate
            best_dist = dist
    return best


def horizontal_overlap_ratio(a_bounds, b_bounds):
    a_left, _, a_width, _ = a_bounds
    b_left, _, b_width, _ = b_bounds
    a_right = a_left + a_width
    b_right = b_left + b_width
    overlap = max(0.0, min(a_right, b_right) - max(a_left, b_left))
    denom = min(a_width, b_width)
    if denom <= 0:
        return 0.0
    return overlap / denom


def find_best_shape(slide, marker_shape, debug=None):
    marker_bounds = shape_bounds(marker_shape)
    marker_left, marker_top, marker_width, marker_height = marker_bounds
    marker_bottom = marker_top + marker_height

    best = None
    best_score = None

    count = slide.Shapes.Count
    for index in range(1, count + 1):
        candidate = slide.Shapes(index)
        if candidate is marker_shape:
            continue

        bounds = shape_bounds(candidate)
        left, top, width, height = bounds
        if width <= 0 or height <= 0:
            continue

        if has_text(candidate):
            continue

        overlap_ratio = horizontal_overlap_ratio(bounds, marker_bounds)
        above = (top + height) <= (marker_top + (marker_height * 0.5))
        if overlap_ratio < 0.25 and not above:
            continue

        area = width * height
        score = area * (1.0 + overlap_ratio)
        if above:
            score += area * 0.25
        if is_picture(candidate):
            score += area
        if is_picture_fill(candidate):
            score += area * 0.5

        if best is None or score > best_score:
            best = candidate
            best_score = score

    if best is not None and debug is not None:
        debug("best placeholder by heuristic found")

    return best


def picture_candidates(shape):
    if is_picture(shape):
        return [shape]

    if is_group(shape):
        candidates = []
        try:
            count = shape.GroupItems.Count
        except Exception:
            return candidates
        for index in range(1, count + 1):
            try:
                item = shape.GroupItems(index)
            except Exception:
                continue
            candidates.extend(picture_candidates(item))
        return candidates

    return []


def add_video(slide, video_path, bounds):
    left, top, width, height = bounds
    try:
        return slide.Shapes.AddMediaObject2(video_path, False, True, left, top, width, height)
    except Exception:
        return slide.Shapes.AddMovie(video_path, False, True, left, top, width, height)


def add_video_url(slide, video_url, bounds):
    left, top, width, height = bounds
    try:
        return slide.Shapes.AddMediaObject2(video_url, True, False, left, top, width, height)
    except Exception:
        return slide.Shapes.AddMovie(video_url, True, False, left, top, width, height)


def apply_bounds(shape, bounds):
    left, top, width, height = bounds
    try:
        shape.LockAspectRatio = 0
    except Exception:
        pass
    try:
        shape.Left = left
        shape.Top = top
        shape.Width = width
        shape.Height = height
    except Exception:
        pass


def apply_rotation(shape, rotation):
    try:
        shape.Rotation = rotation
    except Exception:
        pass


def apply_vertical_align(shape, valign):
    if not valign:
        return
    try:
        if not shape.HasTextFrame:
            return
    except Exception:
        return

    value = str(valign).strip().lower()
    if value in {"top", "start", "baseline"}:
        anchor = 1
    elif value in {"middle", "center", "centre"}:
        anchor = 3
    elif value in {"bottom", "end"}:
        anchor = 4
    else:
        return

    try:
        shape.TextFrame.VerticalAnchor = anchor
    except Exception:
        pass


def is_url(value):
    try:
        parsed = urlparse(value)
    except Exception:
        return False
    return parsed.scheme in {"http", "https"}


def wait_for_pptx_ready(path, retries=10, delay=0.5, debug=None):
    for attempt in range(1, retries + 1):
        try:
            if not os.path.exists(path):
                raise FileNotFoundError(path)
            if os.path.getsize(path) == 0:
                raise OSError("PPTX file is empty.")
            with zipfile.ZipFile(path, "r") as archive:
                if "[Content_Types].xml" not in archive.namelist():
                    raise zipfile.BadZipFile("Missing [Content_Types].xml")
            return True
        except Exception as exc:
            if debug is not None:
                debug(f"pptx not ready (attempt {attempt}/{retries}): {exc}")
            time.sleep(delay)
    return False


def parse_size(value, axis_size=None):
    if value is None:
        return None
    text = str(value).strip().lower()
    if not text:
        return None

    try:
        if text.endswith("%"):
            if axis_size is None:
                return None
            percent = float(text[:-1])
            return axis_size * (percent / 100.0)
        if text.endswith("in"):
            return float(text[:-2]) * 72.0
        if text.endswith("cm"):
            return float(text[:-2]) * (72.0 / 2.54)
        if text.endswith("mm"):
            return float(text[:-2]) * (72.0 / 25.4)
        if text.endswith("pt"):
            return float(text[:-2])
        if text.endswith("px"):
            return float(text[:-2]) * (72.0 / 96.0)
        return float(text)
    except ValueError:
        return None


def apply_custom_size(bounds, width, height, slide_width, slide_height, pos_x, pos_y):
    left, top, base_width, base_height = bounds
    custom_width = parse_size(width, slide_width)
    custom_height = parse_size(height, slide_height)

    if custom_width is None and custom_height is None:
        custom_width = base_width
        custom_height = base_height
    elif custom_width is None:
        aspect = base_width / base_height if base_height else 1.0
        custom_width = custom_height * aspect
    elif custom_height is None:
        aspect = base_height / base_width if base_width else 1.0
        custom_height = custom_width * aspect

    center_x = left + (base_width / 2.0)
    center_y = top + (base_height / 2.0)
    new_left = center_x - (custom_width / 2.0)
    new_top = center_y - (custom_height / 2.0)

    override_x = parse_size(pos_x, slide_width)
    override_y = parse_size(pos_y, slide_height)
    if override_x is not None:
        new_left = override_x
    if override_y is not None:
        new_top = override_y
    return (new_left, new_top, custom_width, custom_height)


def extract_marker(text, prefix):
    if not text:
        return ""
    pattern = re.compile(r"(?<!\\)" + re.escape(prefix) + r"([^\s]+)")
    match = pattern.search(text)
    if not match:
        return ""
    return match.group(1)


def extract_all_markers(text, prefix):
    if not text:
        return []
    pattern = re.compile(r"(?<!\\)" + re.escape(prefix) + r"([^\s]+)")
    return pattern.findall(text)


def extract_layout_marker(text):
    if not text:
        return ""
    match = re.search(r"(?<!\\)" + re.escape(LAYOUT_PREFIX), text)
    if not match:
        return ""
    raw = text[match.end() :]
    first_line = raw.splitlines()[0] if raw else ""
    value = first_line.strip()
    if len(value) >= 2 and value[0] == '"' and value[-1] == '"':
        value = value[1:-1].strip()
    return value


def strip_layout_marker(text):
    if not text or LAYOUT_PREFIX not in text:
        return text
    cleaned = re.sub(r"(?<!\\)" + re.escape(LAYOUT_PREFIX) + r"[^\r\n]*", "", text)
    cleaned = re.sub(r"^[\s\r\n]+", "", cleaned)
    cleaned = re.sub(r"[ \t]+\r", "\r", cleaned)
    cleaned = re.sub(r"\r{2,}", "\r", cleaned)
    return cleaned.strip(" \r\n\t")


def normalize_text(text):
    if text is None:
        return ""
    text = text.replace("\r\n", "\r")
    text = re.sub(r"[ \t]+\r", "\r", text)
    text = re.sub(r"\r{2,}", "\r", text)
    return text.strip(" \r\n\t")


def extract_box_segment(text, box_id):
    start_token = f"{BOX_PREFIX}{box_id}"
    end_token = f"{BOX_END_PREFIX}{box_id}"
    start_index = text.find(start_token)
    if start_index == -1:
        return None
    end_index = text.find(end_token, start_index + len(start_token))
    if end_index == -1:
        return None
    content_start = start_index + len(start_token)
    content_end = end_index
    pre = text[:start_index]
    mid = text[content_start:content_end]
    post = text[end_index + len(end_token) :]
    return pre, mid, post


def strip_caption_markers(text):
    if not text or CAPTION_PREFIX not in text:
        return text
    cleaned = re.sub(r"(?<!\\)" + re.escape(CAPTION_PREFIX) + r"[^\s]+", "", text)
    cleaned = re.sub(r"^[\s\r\n]+", "", cleaned)
    cleaned = re.sub(r"[ \t]+\r", "\r", cleaned)
    cleaned = re.sub(r"\r{2,}", "\r", cleaned)
    return cleaned.strip(" \r\n\t")


def unescape_marker_prefixes(text):
    if not text or "\\" not in text:
        return text
    cleaned = text
    for prefix in ESCAPABLE_MARKER_PREFIXES:
        cleaned = cleaned.replace(f"\\{prefix}", prefix)
    return cleaned


def remove_marker_text(shape, prefix, marker_id):
    if not marker_id:
        return
    try:
        if not shape.HasTextFrame or not shape.TextFrame.HasText:
            return
    except Exception:
        return

    try:
        text_range = shape.TextFrame.TextRange
        text = text_range.Text or ""
    except Exception:
        return

    token = f"{prefix}{marker_id}"
    if token not in text:
        return
    new_text = text.replace(token, "", 1)
    new_text = re.sub(r"^[\s\r\n]+", "", new_text)
    try:
        text_range.Text = new_text
    except Exception:
        pass


def collect_text_shapes(slide, slide_height=None):
    items = []
    for shape in iter_shapes(slide):
        try:
            if shape.HasTextFrame and shape.TextFrame.HasText:
                if is_footer_or_date(shape):
                    continue
                if slide_height is not None:
                    try:
                        if shape.Top > slide_height * 0.85:
                            continue
                    except Exception:
                        pass
                items.append(shape)
        except Exception:
            continue
    return items


def is_footer_or_date(shape):
    try:
        p_type = placeholder_type(shape)
    except Exception:
        p_type = None

    if p_type in {15, 16}:
        return True

    try:
        name = (shape.Name or "").lower()
    except Exception:
        name = ""
    return "footer" in name or "date" in name


def collect_visual_shapes(slide):
    items = []
    for shape in iter_shapes(slide):
        if is_visual_shape(shape):
            items.append(shape)
    return items


def collect_caption_shapes(slide, slide_height=None):
    by_id = {}
    ordered_ids = []
    for shape in collect_text_shapes(slide, slide_height):
        text = shape_text(shape)
        if not text or CAPTION_PREFIX not in text:
            continue
        marker_ids = extract_all_markers(text, CAPTION_PREFIX)
        if not marker_ids:
            continue
        for marker_id in marker_ids:
            if not marker_id or marker_id in by_id:
                continue
            by_id[marker_id] = shape
            ordered_ids.append(marker_id)
    return by_id, ordered_ids


def shape_identity(shape):
    try:
        return shape.Id
    except Exception:
        try:
            return shape.Name
        except Exception:
            return id(shape)


def find_caption_for_image(image_shape, text_shapes, used_ids):
    img_bounds = shape_bounds(image_shape)
    _, img_top, img_width, img_height = img_bounds
    img_bottom = img_top + img_height

    best = None
    best_gap = None

    for shape in text_shapes:
        if shape_identity(shape) in used_ids:
            continue
        bounds = shape_bounds(shape)
        _, top, width, height = bounds
        if width <= 0 or height <= 0:
            continue
        if top < img_top:
            continue

        overlap = horizontal_overlap_ratio(bounds, img_bounds)
        if overlap < 0.25:
            continue

        if height > img_height * 0.6:
            continue
        if width > img_width * 1.2:
            continue

        gap = top - img_bottom
        if gap < -5:
            continue

        if best is None or gap < best_gap:
            best = shape
            best_gap = gap

    if best is not None:
        return best

    # Fallback: caption slightly above the visual.
    best = None
    best_gap = None
    for shape in text_shapes:
        if shape_identity(shape) in used_ids:
            continue
        bounds = shape_bounds(shape)
        _, top, width, height = bounds
        if width <= 0 or height <= 0:
            continue

        overlap = horizontal_overlap_ratio(bounds, img_bounds)
        if overlap < 0.25:
            continue

        if height > img_height * 0.6:
            continue
        if width > img_width * 1.2:
            continue

        gap = img_top - (top + height)
        if gap < -5:
            continue

        if best is None or gap < best_gap:
            best = shape
            best_gap = gap

    return best


def merge_caption_settings(defaults, overrides):
    settings = {}
    settings.update(defaults or {})
    settings.update({k: v for k, v in (overrides or {}).items() if v})
    return settings


def parse_font_size(value):
    if value is None:
        return None
    if str(value).strip().endswith("%"):
        return None
    return parse_size(value, None)


def vertical_overlap_ratio(a_bounds, b_bounds):
    _, a_top, _, a_height = a_bounds
    _, b_top, _, b_height = b_bounds
    a_bottom = a_top + a_height
    b_bottom = b_top + b_height
    overlap = max(0.0, min(a_bottom, b_bottom) - max(a_top, b_top))
    denom = min(a_height, b_height)
    if denom <= 0:
        return 0.0
    return overlap / denom


def group_by_vertical_overlap(items):
    groups = []
    for item in sorted(items, key=lambda x: x["bounds"][1]):
        placed = False
        for group in groups:
            if vertical_overlap_ratio(item["bounds"], group["bounds"]) > 0.5:
                group["items"].append(item)
                group["bounds"] = merge_bounds(group["bounds"], item["bounds"])
                placed = True
                break
        if not placed:
            groups.append({"items": [item], "bounds": item["bounds"]})
    return groups


def merge_bounds(a_bounds, b_bounds):
    a_left, a_top, a_width, a_height = a_bounds
    b_left, b_top, b_width, b_height = b_bounds
    left = min(a_left, b_left)
    top = min(a_top, b_top)
    right = max(a_left + a_width, b_left + b_width)
    bottom = max(a_top + a_height, b_top + b_height)
    return (left, top, right - left, bottom - top)


def placeholder_type(shape):
    try:
        return shape.PlaceholderFormat.Type
    except Exception:
        return None


def collect_slide_layout_rules(entries):
    rules = {}
    if not isinstance(entries, list):
        return rules

    for entry in entries:
        if not isinstance(entry, dict):
            continue
        layout_name = str(entry.get("layout", "")).strip()
        if not layout_name:
            continue
        slide_ref = entry.get("slide", None)
        try:
            slide_index = int(str(slide_ref).strip())
        except Exception:
            continue
        if slide_index <= 0:
            continue
        rules[slide_index] = layout_name
    return rules


def find_slide_layout_marker(slide):
    for shape in iter_shapes(slide):
        alt_text = shape_alt_text(shape)
        name = extract_layout_marker(alt_text)
        if name:
            return name, shape, "alt"

        text_value = shape_text(shape)
        name = extract_layout_marker(text_value)
        if name:
            return name, shape, "text"
    return "", None, ""


def clear_slide_layout_marker(shape, source, debug=None):
    if shape is None:
        return
    try:
        if source == "alt":
            current = shape.AlternativeText or ""
            cleaned = strip_layout_marker(current)
            shape.AlternativeText = cleaned
            return

        if source == "text":
            if not shape.HasTextFrame:
                return
            current = shape.TextFrame.TextRange.Text or ""
            cleaned = strip_layout_marker(current)
            shape.TextFrame.TextRange.Text = cleaned
    except Exception as exc:
        if debug is not None:
            debug(f"failed to clear layout marker: {exc}")


def list_slide_layout_names(slide):
    names = []
    try:
        custom_layouts = slide.Design.SlideMaster.CustomLayouts
    except Exception:
        return names
    for index in range(1, custom_layouts.Count + 1):
        try:
            layout = custom_layouts(index)
            names.append(layout.Name or "")
        except Exception:
            continue
    return names


def apply_slide_layout(slide, layout_name, debug=None):
    wanted = (layout_name or "").strip()
    if not wanted:
        return False
    wanted_lower = wanted.lower()

    try:
        custom_layouts = slide.Design.SlideMaster.CustomLayouts
    except Exception as exc:
        if debug is not None:
            debug(f"custom layouts unavailable on slide {slide.SlideIndex}: {exc}")
        return False

    match = None
    for index in range(1, custom_layouts.Count + 1):
        try:
            layout = custom_layouts(index)
            name = (layout.Name or "").strip()
        except Exception:
            continue
        if name.lower() == wanted_lower:
            match = layout
            break

    if match is None:
        for index in range(1, custom_layouts.Count + 1):
            try:
                layout = custom_layouts(index)
                name = (layout.Name or "").strip()
            except Exception:
                continue
            if wanted_lower in name.lower():
                match = layout
                break

    if match is None:
        if debug is not None:
            available = [name for name in list_slide_layout_names(slide) if name]
            debug(
                f"layout '{wanted}' not found on slide {slide.SlideIndex}. "
                f"available={available}"
            )
        return False

    try:
        slide.CustomLayout = match
        if debug is not None:
            debug(f"applied layout '{match.Name}' on slide {slide.SlideIndex}")
        return True
    except Exception as exc:
        if debug is not None:
            debug(f"failed to apply layout '{wanted}' on slide {slide.SlideIndex}: {exc}")
        return False


def apply_footer_location(presentation, footer_text, location_text, debug=None):
    if not footer_text and not location_text:
        return 0

    footer_type = 15
    date_type = 16
    updated = 0

    def matches_name(shape, token):
        try:
            name = (shape.Name or "").lower()
        except Exception:
            return False
        return token in name

    for slide in presentation.Slides:
        for shape in iter_shapes(slide):
            p_type = placeholder_type(shape)
            if footer_text and (p_type == footer_type or matches_name(shape, "footer")):
                try:
                    if shape.HasTextFrame:
                        shape.TextFrame.TextRange.Text = footer_text
                        updated += 1
                except Exception:
                    pass

            if location_text and (p_type == date_type or matches_name(shape, "date")):
                try:
                    if not shape.HasTextFrame:
                        continue
                    current = shape.TextFrame.TextRange.Text or ""
                    if location_text in current:
                        continue
                    if current.strip():
                        new_text = f"{location_text}, {current.strip()}"
                    else:
                        new_text = location_text
                    shape.TextFrame.TextRange.Text = new_text
                    updated += 1
                except Exception:
                    pass

    if debug is not None:
        debug(f"footer/location placeholders updated: {updated}")
    return updated


def safe_delete(shape, debug=None):
    if shape is None:
        return False
    try:
        shape.Delete()
        return True
    except Exception as exc:
        if debug is not None:
            debug(f"delete failed: {exc}")
        return False


def detect_backend(name):
    value = (name or "auto").strip().lower()
    if value not in {"auto", "msoffice", "python-pptx", "libreoffice", "openoffice"}:
        raise ValueError(f"Unsupported backend: {name}")

    if value == "auto":
        if sys.platform.startswith("win"):
            return "msoffice"
        return "python-pptx"
    if value in {"libreoffice", "openoffice"}:
        return "python-pptx"
    return value


def to_points(emu_value):
    return float(emu_value) / EMU_PER_POINT


def to_emu(points_value):
    return int(round(float(points_value) * EMU_PER_POINT))


def fallback_alt_text(shape):
    try:
        return shape._element.nvSpPr.cNvPr.get("descr", "") or ""
    except Exception:
        return ""


def fallback_set_alt_text(shape, text):
    try:
        shape._element.nvSpPr.cNvPr.set("descr", text)
    except Exception:
        pass


def fallback_shape_text(shape):
    try:
        if shape.has_text_frame:
            return shape.text or ""
    except Exception:
        return ""
    return ""


def fallback_remove_shape(shape):
    try:
        element = shape._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)
            return True
    except Exception:
        return False
    return False


def ensure_soffice_available(debug=None):
    candidates = ("soffice", "openoffice4", "openoffice")
    for name in candidates:
        path = shutil.which(name)
        if path:
            if debug is not None:
                debug(f"office binary detected: {path}")
            return path
    if debug is not None:
        debug("no LibreOffice/OpenOffice binary detected in PATH")
    return ""


def process_with_python_pptx(
    input_path,
    output_path,
    mapping,
    args,
    debug,
    columns_by_id,
    boxes_by_id,
    caption_entries,
    caption_defaults,
    footer_text,
    location_text,
    slide_layout_rules,
):
    del columns_by_id, boxes_by_id, caption_entries, caption_defaults, footer_text, location_text

    try:
        from pptx import Presentation  # type: ignore
    except ImportError as exc:
        print(
            "python-pptx is required on non-Windows platforms. "
            "Install it with: pip install python-pptx"
        )
        raise exc

    # Disclaimer:
    print("Warning: python-pptx backend has limited support, "
          "if your presentation relies on postprocessing features "
          "(layout adjustments, captions, footer/location) it may not work as "
          "expected. For best results, use the msoffice backend on Windows or "
          "post-process it manually.")

    if args.backend in {"libreoffice", "openoffice"}:
        ensure_soffice_available(debug=debug)

    presentation = Presentation(input_path)
    replaced = 0
    missing = []
    skipped_remote = []
    layout_requests = set()

    for slide_index, slide in enumerate(presentation.slides, start=1):
        requested_layout = slide_layout_rules.get(slide_index, "")
        marker_layout = ""
        for shape in slide.shapes:
            alt_marker = extract_layout_marker(fallback_alt_text(shape))
            text_marker = extract_layout_marker(fallback_shape_text(shape))
            found = alt_marker or text_marker
            if found and not marker_layout:
                marker_layout = found
        if requested_layout:
            layout_requests.add(requested_layout)
        if marker_layout:
            layout_requests.add(marker_layout)

        # Iterate on a snapshot because we may delete shapes while processing.
        for shape in list(slide.shapes):
            marker = fallback_alt_text(shape)
            text_value = fallback_shape_text(shape)
            if not marker and PLACEHOLDER_PREFIX in text_value:
                marker = text_value
            if not marker.startswith(PLACEHOLDER_PREFIX):
                continue

            video_id = marker[len(PLACEHOLDER_PREFIX) :].strip()
            item = mapping.get(video_id)
            if not item:
                missing.append(video_id)
                debug(f"missing mapping for id={video_id}")
                continue

            video_path = item.get("video", "")
            if not video_path:
                missing.append(video_id)
                debug(f"missing video file for id={video_id}")
                continue
            if is_url(video_path):
                skipped_remote.append(video_id)
                debug(f"skipping remote URL for id={video_id} on python-pptx backend")
                continue
            if not os.path.exists(video_path):
                missing.append(video_id)
                debug(f"missing video file for id={video_id} path={video_path}")
                continue

            bounds_pt = (
                to_points(shape.left),
                to_points(shape.top),
                to_points(shape.width),
                to_points(shape.height),
            )
            bounds_pt = apply_custom_size(
                bounds_pt,
                item.get("width", ""),
                item.get("height", ""),
                to_points(presentation.slide_width),
                to_points(presentation.slide_height),
                item.get("x", ""),
                item.get("y", ""),
            )

            try:
                movie = slide.shapes.add_movie(
                    video_path,
                    to_emu(bounds_pt[0]),
                    to_emu(bounds_pt[1]),
                    to_emu(bounds_pt[2]),
                    to_emu(bounds_pt[3]),
                    mime_type="video/mp4",
                )
                fallback_set_alt_text(movie, marker)
                fallback_remove_shape(shape)
                replaced += 1
            except Exception as exc:
                debug(f"failed to embed id={video_id}: {exc}")
                missing.append(video_id)

    presentation.save(output_path)

    if not args.keep_mapping:
        for path, label in ((os.path.abspath(args.mapping), "mapping"), (os.path.abspath(args.layout), "layout")):
            try:
                os.remove(path)
                debug(f"deleted {label} file: {path}")
            except FileNotFoundError:
                pass
            except OSError as exc:
                debug(f"failed to delete {label} file: {exc}")

    if missing:
        unique_missing = sorted(set(missing))
        print("Warning: missing videos for placeholders:", ", ".join(unique_missing))
    if skipped_remote:
        unique_remote = sorted(set(skipped_remote))
        print(
            "Warning: remote video URLs are not supported by python-pptx backend: "
            + ", ".join(unique_remote)
        )
    if layout_requests:
        print(
            "Warning: forcing slide layouts is only supported by the Windows msoffice backend. "
            f"Requested layout(s): {', '.join(sorted(layout_requests))}"
        )

    print(f"Embedded {replaced} video(s).")
    if replaced >= 0:
        print("Layout adjustments (columns/boxes/captions/footer) require the Windows msoffice backend.")
    return 0


def main():
    parser = argparse.ArgumentParser(description="Embed videos into a PPTX using placeholders.")
    parser.add_argument("--input", required=True, help="Path to the PPTX to post-process.")
    parser.add_argument("--mapping", default="embed-video.json", help="Path to JSON mapping file.")
    parser.add_argument("--layout", default="pptx-layout.json", help="Path to layout JSON file.")
    parser.add_argument("--output", default="", help="Optional output PPTX path.")
    parser.add_argument("--debug", action="store_true", help="Enable debug logging.")
    parser.add_argument("--keep-mapping", action="store_true", help="Do not delete mapping file.")
    parser.add_argument(
        "--backend",
        default="auto",
        choices=["auto", "msoffice", "python-pptx", "libreoffice", "openoffice"],
        help=(
            "Processing backend. 'auto' uses msoffice on Windows and python-pptx elsewhere. "
            "'libreoffice' and 'openoffice' aliases currently use the python-pptx backend."
        ),
    )
    args = parser.parse_args()

    def debug(msg):
        if debug is not None:
            print(f"[postprocess] {msg}", file=sys.stderr)

    input_path = os.path.abspath(args.input)
    mapping_path = os.path.abspath(args.mapping)
    layout_path = os.path.abspath(args.layout)

    if not os.path.exists(input_path):
        raise FileNotFoundError(f"PPTX not found: {input_path}")

    mapping = {}
    if os.path.exists(mapping_path):
        entries = load_mapping(mapping_path)
        mapping = collect_mapping(entries, os.path.dirname(mapping_path))
        debug(f"mapping file: {mapping_path}")
        debug(f"mapping entries: {len(mapping)}")
        for key, value in mapping.items():
            debug(f"  id={key} video={value.get('video', '')} poster={value.get('poster', '')}")
    else:
        debug(f"mapping file not found: {mapping_path}")

    layout = {"columns": [], "boxes": []}
    if os.path.exists(layout_path):
        layout = load_layout(layout_path)
        debug(f"layout file: {layout_path}")
        debug(f"layout columns: {len(layout.get('columns', []))} layout boxes: {len(layout.get('boxes', []))}")
    else:
        debug(f"layout file not found: {layout_path}")

    columns_by_id, boxes_by_id, caption_entries = collect_layout(layout)
    slide_layout_rules = collect_slide_layout_rules(layout.get("slide_layouts", []))
    footer_text = layout.get("footer", "")
    location_text = layout.get("location", "")
    caption_defaults = layout.get("caption_defaults", {})
    output_path = os.path.abspath(args.output) if args.output else input_path
    backend = detect_backend(args.backend)

    if backend != "msoffice":
        return process_with_python_pptx(
            input_path=input_path,
            output_path=output_path,
            mapping=mapping,
            args=args,
            debug=debug,
            columns_by_id=columns_by_id,
            boxes_by_id=boxes_by_id,
            caption_entries=caption_entries,
            caption_defaults=caption_defaults,
            footer_text=footer_text,
            location_text=location_text,
            slide_layout_rules=slide_layout_rules,
        )

    if not sys.platform.startswith("win"):
        print("msoffice backend is only available on Windows. Use --backend python-pptx on this OS.")
        return 2

    try:
        import win32com.client  # type: ignore
    except ImportError as exc:
        print("pywin32 is required for msoffice backend. Install it with: pip install pywin32")
        raise exc

    if not wait_for_pptx_ready(input_path, debug=debug):
        print("[postprocess] PPTX is not readable yet (invalid zip or still being written).")
        print("[postprocess] Try re-running after the render finishes, or open the PPTX in PowerPoint once to repair it.")
        return 1

    app = win32com.client.Dispatch("PowerPoint.Application")
    app.DisplayAlerts = 0
    presentation = None
    replaced = 0
    adjusted_columns = 0
    adjusted_boxes = 0
    adjusted_captions = 0
    missing = []
    try:
        try:
            presentation = app.Presentations.Open(input_path, WithWindow=False, ReadOnly=False)
        except Exception as exc:
            debug(f"open failed WithWindow=False: {exc}")
            presentation = app.Presentations.Open(input_path, WithWindow=True, ReadOnly=False)
        debug(f"slides: {presentation.Slides.Count}")

        apply_footer_location(presentation, footer_text, location_text, debug=debug)

        for slide in presentation.Slides:
            debug(f"slide {slide.SlideIndex}: shapes={slide.Shapes.Count}")
            rule_layout_name = slide_layout_rules.get(slide.SlideIndex, "")
            marker_layout_name, marker_shape, marker_source = find_slide_layout_marker(slide)
            selected_layout_name = rule_layout_name or marker_layout_name
            if selected_layout_name:
                apply_slide_layout(slide, selected_layout_name, debug=debug)
                if marker_shape is not None:
                    clear_slide_layout_marker(marker_shape, marker_source, debug=debug)

            column_items = []
            box_items = []

            for shape in iter_shapes(slide):
                alt_text = shape_alt_text(shape)
                text_value = shape_text(shape)
                marker = alt_text
                if not marker and PLACEHOLDER_PREFIX in text_value:
                    marker = text_value

                if marker.startswith(PLACEHOLDER_PREFIX):
                    debug(
                        f"found placeholder shape on slide {slide.SlideIndex}: "
                        f"alt_text='{alt_text}' text='{text_value}'"
                    )

                    video_id = marker[len(PLACEHOLDER_PREFIX) :].strip()
                    item = mapping.get(video_id)
                    if not item:
                        missing.append(video_id)
                        debug(f"missing mapping for id={video_id}")
                        continue

                    video_path = item.get("video", "")
                    is_remote = is_url(video_path)
                    if not video_path:
                        missing.append(video_id)
                        debug(f"missing video file for id={video_id} path={video_path}")
                        continue
                    if not is_remote and not os.path.exists(video_path):
                        missing.append(video_id)
                        debug(f"missing video file for id={video_id} path={video_path}")
                        continue

                    target_shape = shape
                    extra_delete = None
                    if not is_picture(shape):
                        best = find_best_shape(slide, shape, debug=debug)
                        if best is None:
                            best = find_nearest_picture(slide, shape)
                        if best is not None:
                            target_shape = best
                            extra_delete = shape
                            debug("marker is text; using heuristic placeholder bounds")

                    bounds = shape_bounds(target_shape)
                    bounds = apply_custom_size(
                        bounds,
                        item.get("width", ""),
                        item.get("height", ""),
                        presentation.PageSetup.SlideWidth,
                        presentation.PageSetup.SlideHeight,
                        item.get("x", ""),
                        item.get("y", ""),
                    )
                    try:
                        rotation = target_shape.Rotation
                    except Exception:
                        rotation = 0.0
                    debug(
                        f"embedding id={video_id} video={video_path} "
                        f"bounds={bounds}"
                    )

                    if is_remote:
                        new_shape = add_video_url(slide, video_path, bounds)
                    else:
                        new_shape = add_video(slide, video_path, bounds)
                    apply_bounds(new_shape, bounds)
                    apply_rotation(new_shape, rotation)
                    try:
                        new_shape.AlternativeText = marker
                    except Exception:
                        pass

                    safe_delete(target_shape, debug=debug)
                    if extra_delete is not None and extra_delete is not target_shape:
                        safe_delete(extra_delete, debug=debug)
                    replaced += 1
                    continue

                if text_value:
                    column_id = extract_marker(text_value, COLUMN_PREFIX)
                    if column_id:
                        entry = columns_by_id.get(column_id, {})
                        column_items.append(
                            {
                                "id": column_id,
                                "shape": shape,
                                "entry": entry,
                                "bounds": shape_bounds(shape),
                            }
                        )
                        continue

                    box_ids = extract_all_markers(text_value, BOX_PREFIX)
                    if box_ids:
                        for box_id in box_ids:
                            entry = boxes_by_id.get(box_id, {})
                            box_items.append(
                                {
                                    "id": box_id,
                                    "shape": shape,
                                    "entry": entry,
                                    "bounds": shape_bounds(shape),
                                }
                            )
                        continue

            if column_items:
                groups = group_by_vertical_overlap(column_items)
                for group in groups:
                    items = sorted(group["items"], key=lambda x: x["bounds"][0])
                    group_left, _, group_width, _ = group["bounds"]

                    specified = []
                    unspecified = []
                    for item in items:
                        width_value = item["entry"].get("width", "")
                        width_points = parse_size(width_value, group_width)
                        if width_points is None:
                            unspecified.append(item)
                        else:
                            specified.append((item, width_points))

                    total_specified = sum(width for _, width in specified)
                    remaining = max(0.0, group_width - total_specified)
                    widths = {}

                    if unspecified:
                        if remaining <= 0:
                            remaining = 0.0
                        base_total = sum(item["bounds"][2] for item in unspecified) or len(unspecified)
                        for item in unspecified:
                            share = item["bounds"][2] / base_total if base_total else 1.0 / len(unspecified)
                            widths[item["id"]] = remaining * share
                    elif total_specified > group_width and total_specified > 0:
                        scale = group_width / total_specified
                        specified = [(item, width * scale) for item, width in specified]

                    for item, width in specified:
                        widths[item["id"]] = width

                    current_left = group_left
                    for item in items:
                        target_width = widths.get(item["id"], item["bounds"][2])
                        left, top, _, height = item["bounds"]
                        new_bounds = (current_left, top, target_width, height)
                        apply_bounds(item["shape"], new_bounds)
                        apply_vertical_align(item["shape"], item["entry"].get("valign", ""))
                        remove_marker_text(item["shape"], COLUMN_PREFIX, item["id"])
                        current_left += target_width
                        adjusted_columns += 1

            for item in box_items:
                shape = item["shape"]
                entry = item["entry"]
                text_value = shape_text(shape)
                segment = extract_box_segment(text_value, item["id"])
                if segment is None:
                    debug(f"box markers not found for id={item['id']}")
                    continue

                pre, mid, post = segment
                box_text = normalize_text(mid)
                remaining_text = normalize_text(pre + post)

                bounds = apply_custom_size(
                    item["bounds"],
                    entry.get("width", ""),
                    entry.get("height", ""),
                    presentation.PageSetup.SlideWidth,
                    presentation.PageSetup.SlideHeight,
                    entry.get("x", ""),
                    entry.get("y", ""),
                )

                if remaining_text:
                    try:
                        shape.TextFrame.TextRange.Text = remaining_text
                    except Exception:
                        pass
                    new_shape = slide.Shapes.AddTextbox(
                        TEXTBOX_ORIENTATION, bounds[0], bounds[1], bounds[2], bounds[3]
                    )
                    try:
                        new_shape.TextFrame.TextRange.Text = box_text
                    except Exception:
                        pass
                    apply_vertical_align(new_shape, entry.get("valign", ""))
                else:
                    try:
                        shape.TextFrame.TextRange.Text = box_text
                    except Exception:
                        pass
                    apply_bounds(shape, bounds)
                    apply_vertical_align(shape, entry.get("valign", ""))

                adjusted_boxes += 1

            if caption_entries:
                caption_shapes_by_id, caption_ids_ordered = collect_caption_shapes(slide)
                if caption_ids_ordered:
                    caption_entries_by_id = {}
                    fallback_entries = []
                    for entry in caption_entries:
                        if not isinstance(entry, dict):
                            continue
                        entry_id = str(entry.get("id", "")).strip()
                        if entry_id:
                            caption_entries_by_id[entry_id] = entry
                        else:
                            fallback_entries.append(entry)

                    used_text_ids = set()
                    fallback_index = 0

                    for caption_id in caption_ids_ordered:
                        caption_shape = caption_shapes_by_id.get(caption_id)
                        if caption_shape is None:
                            continue
                        shape_id = shape_identity(caption_shape)
                        if shape_id in used_text_ids:
                            continue

                        entry = caption_entries_by_id.get(caption_id)
                        if entry is None and fallback_index < len(fallback_entries):
                            entry = fallback_entries[fallback_index]
                            fallback_index += 1
                        if entry is None:
                            entry = {}

                        settings = merge_caption_settings(caption_defaults, entry)

                        font_size = parse_font_size(settings.get("size", ""))
                        if font_size is not None:
                            try:
                                caption_shape.TextFrame.TextRange.Font.Size = font_size
                            except Exception:
                                pass

                        dx = parse_size(settings.get("dx", ""), presentation.PageSetup.SlideWidth)
                        dy = parse_size(settings.get("dy", ""), presentation.PageSetup.SlideHeight)
                        if dx is not None:
                            try:
                                caption_shape.Left = caption_shape.Left + dx
                            except Exception:
                                pass
                        if dy is not None:
                            try:
                                caption_shape.Top = caption_shape.Top + dy
                            except Exception:
                                pass

                        used_text_ids.add(shape_id)
                        adjusted_captions += 1

            # Clean up any leftover caption markers from older runs.
            for shape in collect_text_shapes(slide):
                try:
                    current = shape.TextFrame.TextRange.Text or ""
                except Exception:
                    continue
                if CAPTION_PREFIX not in current:
                    continue
                cleaned = strip_caption_markers(current)
                if cleaned != current:
                    try:
                        shape.TextFrame.TextRange.Text = cleaned
                    except Exception:
                        pass

            # Clean up any leftover layout markers from header attributes.
            for shape in collect_text_shapes(slide):
                try:
                    current = shape.TextFrame.TextRange.Text or ""
                except Exception:
                    continue
                if LAYOUT_PREFIX not in current:
                    continue
                cleaned = strip_layout_marker(current)
                if cleaned != current:
                    try:
                        shape.TextFrame.TextRange.Text = cleaned
                    except Exception:
                        pass

            # Allow writing marker keywords literally with escaped prefixes.
            for shape in collect_text_shapes(slide):
                try:
                    current = shape.TextFrame.TextRange.Text or ""
                except Exception:
                    continue
                if "\\" not in current:
                    continue
                cleaned = unescape_marker_prefixes(current)
                if cleaned != current:
                    try:
                        shape.TextFrame.TextRange.Text = cleaned
                    except Exception:
                        pass

        if output_path != input_path:
            presentation.SaveAs(output_path)
        else:
            presentation.Save()
    finally:
        if presentation is not None:
            presentation.Close()
        app.Quit()

    if not args.keep_mapping:
        for path, label in ((mapping_path, "mapping"), (layout_path, "layout")):
            try:
                os.remove(path)
                debug(f"deleted {label} file: {path}")
            except FileNotFoundError:
                pass
            except OSError as exc:
                debug(f"failed to delete {label} file: {exc}")

    if missing:
        unique_missing = sorted(set(missing))
        print("Warning: missing videos for placeholders:", ", ".join(unique_missing))

    print(f"Embedded {replaced} video(s).")
    if adjusted_columns or adjusted_boxes or adjusted_captions:
        print(
            "Adjusted "
            f"{adjusted_columns} column(s), {adjusted_boxes} box(es), "
            f"and {adjusted_captions} caption(s)."
        )
    return 0


if __name__ == "__main__":
    sys.exit(main())
